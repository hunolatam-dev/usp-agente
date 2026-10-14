# agent/importadores.py — Carga de la cartola BCI
"""
Acepta CSV (separado por ; o ,) y Excel (.xlsx). Las columnas se reconocen por nombre,
sin importar mayusculas, tildes ni el orden, para tolerar los formatos de BCI Pyme.

Clientes, telefonos y facturas NO se importan: se leen directo de app.huno.cl (agent/huno.py).
"""

import csv
import hashlib
import io
import logging
from dataclasses import dataclass, field

from sqlalchemy import select

from agent.db import Movimiento, async_session
from agent.utils import (
    buscar_ruts,
    parsear_fecha,
    parsear_monto,
    sin_tildes,
)

logger = logging.getLogger("usp")

# Alias de columnas (ya sin tildes y en minusculas)
COL_CARTOLA = {
    "fecha": ["fecha", "fecha movimiento", "fecha operacion", "fecha contable"],
    "descripcion": ["descripcion", "glosa", "detalle", "movimiento", "descripcion movimiento"],
    "abono": ["abono", "abonos", "deposito", "depositos", "abonos ($)", "ingreso", "ingresos", "monto abono"],
    "cargo": ["cargo", "cargos", "giro", "giros", "cargos ($)", "egreso"],
    "monto": ["monto", "importe", "monto ($)"],
    "documento": ["n documento", "nro documento", "documento", "n de documento", "numero documento",
                  "n operacion", "nro operacion", "numero operacion", "serial", "n doc"],
    "saldo": ["saldo", "saldo diario", "saldo contable", "saldo disponible"],
}
@dataclass
class ResultadoImportacion:
    leidos: int = 0
    nuevos: int = 0
    actualizados: int = 0
    omitidos: int = 0
    errores: list[str] = field(default_factory=list)

    def dict(self):
        return {"leidos": self.leidos, "nuevos": self.nuevos, "actualizados": self.actualizados,
                "omitidos": self.omitidos, "errores": self.errores[:30]}


def _filas(contenido: bytes, nombre: str) -> list[list]:
    if nombre.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
        hoja = wb.worksheets[0]
        return [list(f) for f in hoja.iter_rows(values_only=True)]
    for codif in ("utf-8-sig", "latin-1"):
        try:
            texto = contenido.decode(codif)
            break
        except UnicodeDecodeError:
            continue
    muestra = texto[:4000]
    sep = ";" if muestra.count(";") >= muestra.count(",") else ","
    if muestra.count("\t") > max(muestra.count(";"), muestra.count(",")):
        sep = "\t"
    return [fila for fila in csv.reader(io.StringIO(texto), delimiter=sep)]


def _norm(c) -> str:
    return " ".join(sin_tildes(str(c or "")).replace("°", " ").replace("º", " ").replace(".", " ")
                    .replace("#", " ").split())


def _mapear(cabecera: list, alias: dict) -> dict[str, int]:
    nombres = [_norm(c) for c in cabecera]
    mapa = {}
    for clave, opciones in alias.items():
        for i, n in enumerate(nombres):
            if n in opciones and i not in mapa.values():
                mapa[clave] = i
                break
    return mapa


def _buscar_cabecera(filas: list[list], alias: dict, requeridas: set[str]) -> tuple[int, dict]:
    for i, fila in enumerate(filas[:40]):
        mapa = _mapear(fila, alias)
        if requeridas <= set(mapa):
            return i, mapa
    raise ValueError(
        "No encontre la fila de titulos. Columnas necesarias: " + ", ".join(sorted(requeridas))
    )


def _celda(fila: list, mapa: dict, clave: str):
    i = mapa.get(clave)
    return fila[i] if i is not None and i < len(fila) else None


# ── Cartola ──────────────────────────────────────────────────────────────


async def importar_cartola(contenido: bytes, nombre: str) -> ResultadoImportacion:
    """Importa solo abonos (ingresos). Re-importar la misma cartola no duplica."""
    res = ResultadoImportacion()
    filas = _filas(contenido, nombre)
    try:
        inicio, mapa = _buscar_cabecera(filas, COL_CARTOLA, {"fecha", "descripcion"})
    except ValueError as e:
        res.errores.append(str(e))
        return res
    if "abono" not in mapa and "monto" not in mapa:
        res.errores.append("La cartola no tiene columna de abonos/depositos ni de monto")
        return res

    vistos: dict[str, int] = {}
    nuevos = []
    for fila in filas[inicio + 1:]:
        if not any(c not in (None, "") for c in fila):
            continue
        fecha = parsear_fecha(_celda(fila, mapa, "fecha"))
        if fecha is None:
            continue  # totales, subtitulos, pie de pagina
        res.leidos += 1
        if "abono" in mapa:
            monto = parsear_monto(_celda(fila, mapa, "abono"))
        else:
            monto = parsear_monto(_celda(fila, mapa, "monto"))
            if monto is not None and monto < 0:
                monto = None
        if not monto or monto <= 0:
            res.omitidos += 1
            continue
        desc = str(_celda(fila, mapa, "descripcion") or "").strip()
        doc = str(_celda(fila, mapa, "documento") or "").strip() or None
        saldo = str(_celda(fila, mapa, "saldo") or "").strip()
        base = f"{fecha.isoformat()}|{monto}|{desc}|{doc or ''}|{saldo}"
        vistos[base] = vistos.get(base, 0) + 1
        huella = hashlib.sha256(f"{base}|{vistos[base]}".encode()).hexdigest()
        ruts = buscar_ruts(desc)
        nuevos.append(Movimiento(fecha=fecha, monto=monto, disponible=monto, descripcion=desc[:2000], documento=doc,
                                 rut_origen=ruts[0] if ruts else None, huella=huella, fuente="cartola"))

    async with async_session() as s:
        existentes = set()
        huellas = [m.huella for m in nuevos]
        for i in range(0, len(huellas), 500):
            existentes |= set((await s.execute(select(Movimiento.huella).where(Movimiento.huella.in_(huellas[i:i + 500])))).scalars())
        for m in nuevos:
            if m.huella in existentes:
                res.omitidos += 1
                continue
            s.add(m)
            res.nuevos += 1
        await s.commit()
    return res
